#!/usr/bin/env python3
# SPDX-License-Identifier: CERN-OHL-S-2.0
# Copyright 2026 Mike Matthews.
"""U5 Option A: cascading parametric model workbook -> docs-authority/MODEL_WORKBOOK.xlsx.

Human-readable view of the whole plant calc chain, regenerated with the graph like
every derived file.  No hand-maintained numbers: every cell is read live from the
unified sources --

  * flow_graph/param/drivers.json       696 cited keys (the independent inputs)
  * flow_graph/plant_flow_graph.json    the sole authority (units, edges, mvs block)
  * flow_graph/system_tree.yaml         tree-of-trees; kW rollups via SystemTree.aggregate

Sheets: README (provenance), DRIVERS (all keys sorted by dotted path + tree split),
TREE (levels, parents, unit counts, rolled-up kW), ELECTRICAL (mvs.electrical +
PV/storage/inverter chain), HEAT_LEDGER (Q1-Q4 thermal-kW edge table + SUMIF store
balance; the plain name HEAT belongs to the HEAT sector sheet), and one sheet per
tree sector (19).  Formulas are real where honest -- per-sector kW subtotals,
README counts, HEAT_LEDGER store balance -- everything else is data.

Byte-stable by construction: no wall-clock anywhere.  docProps dates are derived
from the graph sha256 only and every zip entry is rewritten with a fixed timestamp,
so two runs over the same graph produce identical bytes.  Sector sheets are keyed
by tree roots (the 19 mvs sector ids), so a unit whose graph sector_id is a
pre-rename legacy id (WET, sector_exceptions) sits on its owning root's sheet with
its graph sector shown.

Regenerate: python3 -m flow_graph.param.regen            (regen stage `model_workbook`)
            python3 docs-authority/_model_workbook.py    (standalone)
"""
import hashlib
import json
import os
import re
import sys
import zipfile
from datetime import datetime, timedelta

R = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(R)
sys.path.insert(0, R)                       # flow_graph imports for the tree aggregate

GRAPH = 'flow_graph/plant_flow_graph.json'
DRIVERS = 'flow_graph/param/drivers.json'
TREE = 'flow_graph/system_tree.yaml'
OUT = 'docs-authority/MODEL_WORKBOOK.xlsx'

# ---------------------------------------------------------------- style constants
HDR_FILL = 'FF1F3864'      # dark blue header band
HDR_FONT = 'FFFFFFFF'
INPUT_FILL = 'FFFFF2CC'    # pale yellow: driver values (independent inputs)
DERIVED_FILL = 'FFE2EFDA'  # pale green: rolled-up / summed kW
SECTION_FILL = 'FFD9E1F2'  # pale blue: section titles inside a sheet


def _scalar(v):
    """One cell value: numbers/bools/str pass through; containers compact-JSON."""
    if isinstance(v, (int, float, bool)) or v is None:
        return v
    if isinstance(v, str):
        return v
    return json.dumps(v, sort_keys=True, separators=(',', ':'))


def _header(ws, row, cols, widths=None):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = _F('bold', color=HDR_FONT)
        c.fill = _FILL(HDR_FILL)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def _F(style='regular', color=None):
    from openpyxl.styles import Font
    return Font(bold=(style == 'bold'), color=color)


def _FILL(argb):
    from openpyxl.styles import PatternFill
    return PatternFill('solid', fgColor=argb)


def _kv(ws, row, field, value, note=None):
    ws.cell(row=row, column=1, value=field).font = _F('bold')
    ws.cell(row=row, column=2, value=value)
    if note:
        ws.cell(row=row, column=3, value=note)


# ---------------------------------------------------------------- sheet builders

def sheet_readme(wb, gsha, g, tree, stats, sector_ids):
    ws = wb.create_sheet('README')
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 96
    ws.column_dimensions['C'].width = 60
    rev = g.get('revision')
    tstats = tree.stats()
    rows = [
        ('MODEL_WORKBOOK', 'Cascading parametric model workbook (U5, Option A): the plant calc chain as one readable view.'),
        ('STATUS', 'GENERATED FILE -- DO NOT EDIT BY HAND. Regeneration overwrites every sheet.'),
        ('GENERATOR', 'docs-authority/_model_workbook.py'),
        ('REGENERATE', 'python3 -m flow_graph.param.regen  (stage `model_workbook`, after task_plan)   |   standalone: python3 docs-authority/_model_workbook.py'),
        ('GRAPH', f'{GRAPH} -- revision {rev} -- sha256 {gsha}'),
        ('DRIVERS', f'{DRIVERS} -- keys:', '=COUNTA(DRIVERS!A:A)-1'),
        ('SYSTEM TREE', f'{TREE} -- {tree.doc_id} Rev {tree.revision} ({tree.date}) -- nodes:', '=COUNTA(TREE!A:A)-1'),
        ('PLANT SET', f"{stats['units']} units across {len(sector_ids)} sectors; the tree leaf set equals the graph unit set (drift guard: pytest flow_graph/param/test_system_tree.py)."),
        ('SOURCES RULE', 'No hand-maintained numbers: drivers from drivers.json, unit/edge/mvs values from the graph, rollups from SystemTree.aggregate.'),
        ('FORMULAS', 'Live Excel formulas where honest: per-sector SUBTOTAL kW (SUM), this sheet\'s DRIVERS/TREE counts (COUNTA), HEAT_LEDGER store balance (SUMIF). All other cells are values.'),
        ('DATES', 'No wall clock: workbook docProps dates are derived from the graph sha256 only, so the file is byte-stable across runs over the same graph.'),
        ('NAMING NOTE', 'The heat summary sheet is HEAT_LEDGER because the plain name HEAT is the HEAT sector sheet; ELECTRICAL does not collide (the sector is ELEC).'),
    ]
    for i, row in enumerate(rows, 1):
        a, b = row[0], row[1]
        c = row[2] if len(row) > 2 else None
        ws.cell(row=i, column=1, value=a).font = _F('bold')
        ws.cell(row=i, column=2, value=b)
        if c:
            cc = ws.cell(row=i, column=3, value=c)
            cc.fill = _FILL(DERIVED_FILL)
    r = len(rows) + 2
    sec = ws.cell(row=r, column=1, value='SHEET INDEX')
    sec.font = _F('bold')
    r += 1
    _header(ws, r, ['sheet', 'contents'], widths=None)
    index = [
        ('README', 'this provenance page'),
        ('DRIVERS', f"all {stats['drivers']} drivers.json keys sorted by dotted path; value/units/ref/register/needs_verification/verification_test + the tree-path split columns"),
        ('TREE', f"{stats['tree_nodes']} system-tree nodes: level, parent, type, unit counts, rolled-up kW (SystemTree.aggregate)"),
        ('ELECTRICAL', 'mvs.electrical block (need/target/PV/ZnBr2/rectifier) + the PV / storage / inverter chain read from the graph units'),
        ('HEAT_LEDGER', 'Q1-Q4 ledger: every thermal-layer kW edge, then the per-store in/out/net balance (SUMIF formulas)'),
    ] + [(s, f"sector sheet: units (continuous_kw, presence, class, enabled) + SUBTOTAL kW + every edge touching the sector") for s in sector_ids]
    for name, what in index:
        r += 1
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=what)
    return ws


def sheet_drivers(wb, drivers):
    ws = wb.create_sheet('DRIVERS')
    depth = max(k.count('.') for k in drivers) + 1
    cols = ['key', 'value', 'units', 'ref', 'register', 'needs_verification',
            'verification_test'] + [f'tree_{i}' for i in range(1, depth + 1)]
    widths = [46, 22, 20, 26, 30, 18, 46] + [20] * depth
    _header(ws, 1, cols, widths)
    r = 1
    for key in sorted(drivers):
        e = drivers[key]
        r += 1
        ws.cell(row=r, column=1, value=key)
        v = ws.cell(row=r, column=2, value=_scalar(e.get('value')))
        v.fill = _FILL(INPUT_FILL)                      # independent input
        ws.cell(row=r, column=3, value=e.get('units'))
        ws.cell(row=r, column=4, value=e.get('ref') or e.get('owner') or e.get('published') or e.get('basis'))
        ws.cell(row=r, column=5, value=e.get('register'))
        ws.cell(row=r, column=6, value=e.get('needs_verification'))
        ws.cell(row=r, column=7, value=e.get('verification_test'))
        if '.' in key:
            for i, part in enumerate(key.split('.'), 1):
                ws.cell(row=r, column=7 + i, value=part)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(cols)).column_letter}{r}'
    return r


def sheet_tree(wb, tree):
    ws = wb.create_sheet('TREE')
    cols = ['node_id', 'parent', 'level', 'type', 'unit_id', 'description',
            'units', 'kw_sum', 'kw_stamped', 'worst_presence']
    _header(ws, 1, cols, [52, 52, 7, 11, 24, 40, 8, 12, 10, 15])
    r = 1
    for n in sorted(tree.nodes, key=lambda n: n.id):
        agg = tree.aggregate(tree.resolve_units(n.id))
        r += 1
        ws.cell(row=r, column=1, value=n.id)
        ws.cell(row=r, column=2, value=n.parent)
        ws.cell(row=r, column=3, value=n.level)
        kind = ('unit' if n.unit_id is not None else
                'root' if n.level == 0 else 'train' if n.level == 1 else 'subsystem')
        ws.cell(row=r, column=4, value=kind)
        ws.cell(row=r, column=5, value=n.unit_id)
        ws.cell(row=r, column=6, value=n.description)
        ws.cell(row=r, column=7, value=agg['count'])
        k = ws.cell(row=r, column=8, value=agg['kw_sum'])
        k.fill = _FILL(DERIVED_FILL)
        k.number_format = '0.0000'
        ws.cell(row=r, column=9, value=agg['kw_count'])
        ws.cell(row=r, column=10, value=agg['presence'])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:J{r}'
    return r


def sheet_sector(wb, sid, units, node_of, edges):
    ws = wb.create_sheet(sid)
    cols = ['unit_id', 'continuous_kw', 'kw_unit', 'kw_status', 'presence',
            'class', 'enabled', 'graph_sector']
    _header(ws, 1, cols, [24, 14, 8, 10, 12, 16, 9, 13])
    r = 1
    for uid in units:
        n = node_of[uid]
        r += 1
        ws.cell(row=r, column=1, value=uid)
        cw = n.get('continuous_kw')
        if isinstance(cw, dict):
            ws.cell(row=r, column=2, value=cw.get('value')).number_format = '0.0000'
            ws.cell(row=r, column=3, value=cw.get('unit'))
            ws.cell(row=r, column=4, value=cw.get('status'))
        elif isinstance(cw, (int, float)) and not isinstance(cw, bool):
            c = ws.cell(row=r, column=2, value=cw)      # Rev AC shorthand: bare number = kW study
            c.number_format = '0.0000'
            ws.cell(row=r, column=3, value='kW')
        ws.cell(row=r, column=5, value=n.get('presence'))
        ws.cell(row=r, column=6, value=n.get('class'))
        ws.cell(row=r, column=7, value=n.get('enabled'))
        gs = ws.cell(row=r, column=8, value=n.get('sector_id'))
        if n.get('sector_id') != sid:
            gs.fill = _FILL(SECTION_FILL)               # sector_exception unit (e.g. WET under DIGESTER)
    first, last = 2, r
    r += 2
    lab = ws.cell(row=r, column=1, value='SUBTOTAL continuous kW (SUM)')
    lab.font = _F('bold')
    sub = ws.cell(row=r, column=2, value=f'=SUM(B{first}:B{last})')
    sub.font = _F('bold')
    sub.fill = _FILL(DERIVED_FILL)
    sub.number_format = '0.0000'
    ws.auto_filter.ref = f'A1:H{last}'
    ws.freeze_panes = 'A2'
    r += 2
    t = ws.cell(row=r, column=1, value='EDGES TOUCHING THIS SECTOR (from or to)')
    t.font = _F('bold')
    t.fill = _FILL(SECTION_FILL)
    r += 1
    ecols = ['stream_id', 'from_unit', 'to_unit', 'rate', 'unit', 'status']
    _header(ws, r, ecols)
    for e in edges:
        r += 1
        ws.cell(row=r, column=1, value=e['stream_id'])
        ws.cell(row=r, column=2, value=e['from_unit_id'])
        ws.cell(row=r, column=3, value=e['to_unit_id'])
        rate = e.get('rate') or {}
        ws.cell(row=r, column=4, value=rate.get('value'))
        ws.cell(row=r, column=5, value=rate.get('unit'))
        ws.cell(row=r, column=6, value=rate.get('status'))
    return r


ELEC_CHAIN = [  # (unit, param path) -- explicit, so nothing is guessed
    ('PV_ARRAY', 'module_kw'), ('PV_ARRAY', 'count'), ('PV_ARRAY', 'nameplate_kw'),
    ('PV_ARRAY', 'area_m2'), ('PV_ARRAY', 'balance_kWp'), ('PV_ARRAY', 'oversize_factor'),
    ('ZNBR2', 'kW'), ('ZNBR2', 'kWh'), ('ZNBR2', 'charge_kW'), ('ZNBR2', 'discharge_kW'),
    ('SOL_ARK_60K', 'kW'), ('SOL_ARK_60K', 'count'), ('SOL_ARK_60K', 'dc_ac_ratio'),
    ('SOL_ARK_60K', 'interconnection.export_peak_kW'),
    ('RECTIFIER', 'rated_A'), ('RECTIFIER', 'nameplate_kw_dc'),
    ('RECTIFIER', 'rect_ac_input_kW'), ('RECTIFIER', 'rect_loss_kW'),
    ('BACKUP_GEN', 'rated_kW'),
]


def sheet_electrical(wb, g, node_of):
    ws = wb.create_sheet('ELECTRICAL')
    _header(ws, 1, ['block', 'item', 'value', 'unit', 'status', 'note'],
            [14, 34, 14, 9, 10, 90])
    ws.freeze_panes = 'A2'
    el = (g.get('mvs') or {}).get('electrical') or {}
    r = 1
    for item in ('need_kw', 'target_kw'):
        v = el.get(item) or {}
        r += 1
        ws.cell(row=r, column=1, value='mvs.electrical')
        ws.cell(row=r, column=2, value=item)
        ws.cell(row=r, column=3, value=v.get('value')).number_format = '0.00'
        ws.cell(row=r, column=4, value=v.get('unit'))
        ws.cell(row=r, column=5, value=v.get('status'))
        ws.cell(row=r, column=6, value=v.get('note'))
    for item in ('pv_module_kw', 'pv_count', 'znbr_module_kw', 'znbr_count', 'rect_count'):
        r += 1
        ws.cell(row=r, column=1, value='mvs.electrical')
        ws.cell(row=r, column=2, value=item)
        ws.cell(row=r, column=3, value=el.get(item))
    for item, val in (('bus_nodes', ' | '.join(el.get('bus_nodes') or [])),
                      ('storage_options', ' | '.join(el.get('storage_options') or []))):
        r += 1
        ws.cell(row=r, column=1, value='mvs.electrical')
        ws.cell(row=r, column=2, value=item)
        ws.cell(row=r, column=3, value=val)
    r += 1
    from openpyxl.styles import Alignment
    n = ws.cell(row=r, column=1, value=el.get('notes'))
    n.alignment = Alignment(wrap_text=True)
    chain_units = []
    for uid, path in ELEC_CHAIN:
        if uid not in chain_units:
            chain_units.append(uid)
    for uid in chain_units:
        node = node_of.get(uid)
        r += 1
        cw = (node or {}).get('continuous_kw') or {}
        ws.cell(row=r, column=1, value=uid)
        ws.cell(row=r, column=2, value='continuous_kw')
        ws.cell(row=r, column=3, value=cw.get('value')).number_format = '0.0000'
        ws.cell(row=r, column=4, value=cw.get('unit'))
        ws.cell(row=r, column=5, value=cw.get('status'))
    for uid, path in ELEC_CHAIN:
        node = node_of.get(uid) or {}
        p = node.get('params') or {}
        for part in path.split('.'):
            p = p.get(part) if isinstance(p, dict) else None
        if isinstance(p, (int, float)) and not isinstance(p, bool):
            r += 1
            ws.cell(row=r, column=1, value=uid)
            ws.cell(row=r, column=2, value=f'params.{path}')
            ws.cell(row=r, column=3, value=p)
    return r


def sheet_heat(wb, g):
    ws = wb.create_sheet('HEAT_LEDGER')
    _header(ws, 1, ['stream_id', 'from_unit', 'to_unit', 'kw', 'status'],
            [24, 22, 22, 11, 10])
    ws.freeze_panes = 'A2'
    led = [e for e in g['edges']
           if 'thermal' in (e.get('layers') or [])
           and e.get('presence') != 'retired'
           and (e.get('rate') or {}).get('unit') == 'kW']
    r = 1
    for e in sorted(led, key=lambda e: e['stream_id']):
        r += 1
        ws.cell(row=r, column=1, value=e['stream_id'])
        ws.cell(row=r, column=2, value=e['from_unit_id'])
        ws.cell(row=r, column=3, value=e['to_unit_id'])
        k = ws.cell(row=r, column=4, value=(e.get('rate') or {}).get('value'))
        k.number_format = '0.0000'
        ws.cell(row=r, column=5, value=(e.get('rate') or {}).get('status'))
    first, last = 2, r
    units = sorted({e['from_unit_id'] for e in led} | {e['to_unit_id'] for e in led},
                   key=lambda u: (0 if u[:3] in ('Q1_', 'Q2_', 'Q3_', 'Q4_') else 1, u))
    r += 2
    t = ws.cell(row=r, column=1, value='STORE BALANCE (SUMIF over the ledger edges; the X-THM-01 tolerance is 0.01 kW)')
    t.font = _F('bold')
    t.fill = _FILL(SECTION_FILL)
    r += 1
    _header(ws, r, ['unit', 'in_kW', 'out_kW', 'net_kW'])
    rng_from = f'$B${first}:$B${last}'
    rng_to = f'$C${first}:$C${last}'
    rng_kw = f'$D${first}:$D${last}'
    for u in units:
        r += 1
        ws.cell(row=r, column=1, value=u)
        i = ws.cell(row=r, column=2, value=f'=SUMIF({rng_to},A{r},{rng_kw})')
        o = ws.cell(row=r, column=3, value=f'=SUMIF({rng_from},A{r},{rng_kw})')
        n = ws.cell(row=r, column=4, value=f'=B{r}-C{r}')
        for c in (i, o, n):
            c.fill = _FILL(DERIVED_FILL)
            c.number_format = '0.0000'
    return r, len(led), len(units)


# ---------------------------------------------------------------- deterministic save

def _normalise_zip(path, stamp):
    """openpyxl embeds wall-clock zip timestamps AND rewrites docProps/modified at save
    time; rewrite every entry with a fixed timestamp and the stamp-derived core dates so
    two runs over the same graph produce identical bytes."""
    iso = stamp.strftime('%Y-%m-%dT%H:%M:%SZ')
    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, blob in entries:
            if name == 'docProps/core.xml':
                text = blob.decode('utf-8')
                text = re.sub(r'(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)',
                              lambda m: m.group(1) + iso + m.group(2), text)
                blob = text.encode('utf-8')
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, blob)


def build():
    from openpyxl import Workbook
    import yaml
    from flow_graph.system_tree import load as load_tree

    g = json.load(open(GRAPH))
    drivers = json.load(open(DRIVERS))
    tree = load_tree()                                  # raises on a broken/drifted tree
    report = tree.validate()                            # ... and on graph drift
    node_of = {n['unit_id']: n for n in g['nodes']}
    sector_ids = [s['sector_id'] for s in (g.get('mvs') or {}).get('sectors') or []]
    gsha = hashlib.sha256(open(GRAPH, 'rb').read()).hexdigest()

    wb = Workbook()
    wb.remove(wb.active)
    # No wall clock: the workbook date is a function of the graph content only.
    stamp = datetime(2026, 1, 1) + timedelta(days=int(gsha[:8], 16) % 365)
    wb.properties.creator = 'docs-authority/_model_workbook.py (generated from the unified calc sources)'
    wb.properties.lastModifiedBy = 'docs-authority/_model_workbook.py'
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.title = 'AQS model workbook (generated)'
    wb.calculation.fullCalcOnLoad = True                # evaluate the SUM/COUNTA/SUMIF cells on open

    stats = {'drivers': len(drivers), 'tree_nodes': len(tree.nodes),
             'units': report['graph_units'], 'leaves': report['leaves']}
    sheet_readme(wb, gsha, g, tree, stats, sector_ids)
    dr = sheet_drivers(wb, drivers)
    tr = sheet_tree(wb, tree)
    er = sheet_electrical(wb, g, node_of)
    hr, n_led, n_stores = sheet_heat(wb, g)
    sec_rows = {}
    for sid in sector_ids:                              # 19 sheets, keyed by tree roots
        units = sorted(tree.resolve_units(sid))
        edges = sorted((e for e in g['edges']
                        if e.get('from_sector_id') == sid or e.get('to_sector_id') == sid),
                       key=lambda e: e['stream_id'])
        sec_rows[sid] = sheet_sector(wb, sid, units, node_of, edges)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    _normalise_zip(OUT, stamp)
    print(f"MODEL_WORKBOOK.xlsx: {len(wb.sheetnames)} sheets -- DRIVERS {dr - 1} rows, "
          f"TREE {tr - 1} rows, ELECTRICAL {er - 1} rows, HEAT_LEDGER {n_led} edges / {n_stores} stores; "
          f"{len(sector_ids)} sector sheets ({min(sec_rows.values())}-{max(sec_rows.values())} rows); "
          f"graph {gsha[:12]} rev {g.get('revision')}, tree validated ({report['leaves']} leaves = graph units)")


if __name__ == '__main__':
    build()
